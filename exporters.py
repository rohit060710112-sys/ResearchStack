"""
exporters.py
------------
Export utilities to generate PDF and DOCX files from markdown reports.
"""

import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


def export_to_docx(topic: str, markdown_text: str) -> io.BytesIO:
    """Generate a cleanly formatted Word document (.docx) from markdown text."""
    doc = Document()

    # Document Title
    title = doc.add_heading(f"Research Report: {topic}", level=0)
    title.runs[0].font.color.rgb = RGBColor(220, 100, 30)

    for line in markdown_text.splitlines():
        line_clean = line.strip()
        if not line_clean:
            continue

        if line_clean.startswith("# "):
            h = doc.add_heading(line_clean[2:], level=1)
            h.runs[0].font.color.rgb = RGBColor(40, 40, 50)
        elif line_clean.startswith("## "):
            h = doc.add_heading(line_clean[3:], level=2)
            h.runs[0].font.color.rgb = RGBColor(60, 60, 70)
        elif line_clean.startswith("### "):
            h = doc.add_heading(line_clean[4:], level=3)
        elif line_clean.startswith("- ") or line_clean.startswith("* "):
            doc.add_paragraph(line_clean[2:], style="List Bullet")
        elif line_clean.startswith("1. ") or line_clean.startswith("2. ") or line_clean.startswith("3. "):
            doc.add_paragraph(line_clean[3:], style="List Number")
        else:
            p = doc.add_paragraph(line_clean)
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(6)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_to_pdf(topic: str, markdown_text: str) -> io.BytesIO:
    """Generate a styled PDF document from report text using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#d9531e'),
        spaceAfter=14
    )
    
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1f2421'),
        spaceBefore=10,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#2b2b2b'),
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        leftIndent=15,
        textColor=colors.HexColor('#2b2b2b'),
        spaceAfter=4
    )

    story = [
        Paragraph(f"Research Report: {topic}", title_style),
        Spacer(1, 10)
    ]

    for line in markdown_text.splitlines():
        clean = line.strip()
        if not clean:
            continue
        
        # Escape XML entities for ReportLab Paragraph
        safe_line = clean.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

        if safe_line.startswith("# ") or safe_line.startswith("## ") or safe_line.startswith("### "):
            header_text = safe_line.lstrip('#').strip()
            story.append(Paragraph(header_text, h2_style))
        elif safe_line.startswith("- ") or safe_line.startswith("* "):
            bullet_text = "&bull; " + safe_line[2:].strip()
            story.append(Paragraph(bullet_text, bullet_style))
        else:
            story.append(Paragraph(safe_line, body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer

